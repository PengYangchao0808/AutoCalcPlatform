"""Report schema v2: versioned payload, full config, coverage, provenance (todo 24, gap §8.2/§12.1).

BEFORE (captured verbatim in
``.omo/evidence/acp-nmr-goodman-gap-remediation/task-24-before-v1-shape.json``):
the report had no ``schema_version``; ``config`` carried only
``{nmr_method, nmr_basis, solvent}``; assignments had no ``scaled_ppm``;
candidates carried no ``analysis_status``/``coverage``; there was no
calibration/TMS provenance. Legacy (v1) payloads — identified by the absence
of ``schema_version`` — must keep reading unchanged and render
``历史报告：验证状态未知`` (never auto-upgraded).

JSON ↔ XLSX consistency: both artifacts serialize the SAME numbers —
``scaled_ppm`` is the raw field value in both, every other number uses the
same serialization precision on both sides.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from acp.nmr.models import (
    REPORT_SCHEMA_VERSION,
    Assignment,
    CandidateEvidence,
    CandidateProbability,
    CandidateResult,
    ConformerShielding,
    NmrConfig,
    NmrReport,
    NucleusEvidence,
    ProbabilityResult,
    RegressionResult,
)
from acp.nmr.report import (
    LEGACY_REPORT_NOTE,
    report_validation_note,
    write_json_report,
    write_xlsx_report,
)

# raw scaled value — serialized unchanged (never a display rounding)
_RAW_SCALED = 4.100000012345
# non-4-decimal inputs exercise the shared round-4 serialization
_EXP = 4.0000123456
_CALC = 4.123456789
_RESIDUAL = -0.0999987654
# Boltzmann weights of the successful conformers (sum deliberately ≠ 1)
_CONF_WEIGHTS = (0.5, 0.3, 0.14876)

# Legacy (v1) fixture: the exact pre-v2 shape captured BEFORE todo 24.
_LEGACY_V1_FIXTURE = json.dumps(
    {
        "summary": {
            "n_candidates": 1,
            "winner": {"index": 0, "label": "cand_1", "dp4": 0.5, "dp5": None},
            "dp4_ranking": "not_applicable",
            "nuclei": ["1H", "13C"],
        },
        "candidates": [
            {
                "index": 0,
                "label": "cand_1",
                "dp4_probability": 0.5,
                "dp5_probability": None,
                "dp5_diagnostic_score": None,
                "dp5_kernel": None,
                "evidence": None,
                "probability": None,
                "n_conformers": 0,
                "regression": {},
                "assignment": [
                    {
                        "atom": "H1",
                        "element": "H",
                        "exp_ppm": 4.0,
                        "calc_ppm": 4.1235,
                        "residual": -0.1,
                    }
                ],
                "conformers": [],
            }
        ],
        "config": {
            "nmr_method": "mPW1PW91",
            "nmr_basis": "6-311G(d)",
            "solvent": "chloroform",
        },
        "error_model": "goodman-legacy",
        "dp5_mode": "fallback",
        "fchl_kernel": "",
        "note": "",
    },
    indent=2,
    ensure_ascii=False,
)


# ---------------------------------------------------------------------------
# fixtures
# ---------------------------------------------------------------------------


def _assignment() -> Assignment:
    return Assignment(
        atom_label="H1",
        element="H",
        exp_ppm=_EXP,
        calc_ppm=_CALC,
        scaled_ppm=_RAW_SCALED,
        residual=_RESIDUAL,
    )


def _valid_evidence() -> CandidateEvidence:
    return CandidateEvidence(
        status="valid",
        per_nucleus={"13C": NucleusEvidence(1, 1), "1H": NucleusEvidence(3, 3)},
        observation_ids=("13C:0", "1H:0", "1H:1", "1H:2"),
        exclusion_reasons=(),
        total_matched=4,
    )


def _invalid_evidence() -> CandidateEvidence:
    return CandidateEvidence(
        status="invalid",
        per_nucleus={"13C": NucleusEvidence(1, 0), "1H": NucleusEvidence(3, 0)},
        observation_ids=(),
        exclusion_reasons=("no_matched_signals",),
        total_matched=0,
    )


def _probability(
    dp4_status: str,
    dp5_status: str,
    dp4: float | None,
    dp5: float | None,
    dp4_calib: str,
    dp5_calib: str,
) -> CandidateProbability:
    return CandidateProbability(
        dp4=ProbabilityResult(
            model_id="goodman-dp4",
            model_version="goodman-legacy",
            status=dp4_status,  # type: ignore[arg-type]
            probability=dp4,
            mode=None,
            calibration_status=dp4_calib,
        ),
        dp5=ProbabilityResult(
            model_id="goodman-dp5",
            model_version="goodman-legacy",
            status=dp5_status,  # type: ignore[arg-type]
            probability=dp5,
            mode="fchl" if dp5 is not None else None,
            calibration_status=dp5_calib,
        ),
    )


def _conformers() -> list[ConformerShielding]:
    return [
        ConformerShielding(
            conformer_id=f"conf_{i:03d}",
            boltzmann_weight=weight,
            shieldings={0: {"symbol": "C", "isotropic": 188.0}},
        )
        for i, weight in enumerate(_CONF_WEIGHTS)
    ]


def _report() -> NmrReport:
    valid = CandidateResult(
        index=0,
        label="cand_0",
        assignments=[_assignment()],
        regressions={
            "13C": RegressionResult(
                nucleus="13C",
                slope=1.0123456789,
                intercept=-0.5432109876,
                r_squared=0.987654321,
                mae=0.1234567891,
            )
        },
        dp4_probability=0.7184267845659725,
        dp5_probability=0.4910561234567,
        conformer_shieldings=_conformers(),
        evidence=_valid_evidence(),
        probability=_probability(
            "valid",
            "valid",
            0.7184267845659725,
            0.4910561234567,
            "unvalidated_protocol",
            "goodman_kde",
        ),
    )
    excluded = CandidateResult(
        index=1,
        label="cand_1",
        evidence=_invalid_evidence(),
        probability=_probability(
            "invalid",
            "unavailable",
            None,
            None,
            "not_evaluated",
            "not_evaluated",
        ),
    )
    legacy = CandidateResult(index=2, label="legacy")  # no evidence/probability/conformers
    return NmrReport(
        candidates=[valid, excluded, legacy],
        config=NmrConfig(),
        error_model="goodman-legacy",
        dp5_mode="fchl",
        metadata={"fchl_kernel": "qml"},
    )


# ---------------------------------------------------------------------------
# Assignment.as_dict: scaled_ppm added, raw value
# ---------------------------------------------------------------------------


def test_assignment_scaled_ppm_is_raw_value() -> None:
    """scaled_ppm serializes the raw field — not any display rounding."""
    data = _assignment().as_dict()
    # existing keys kept
    assert {"atom", "element", "exp_ppm", "calc_ppm", "residual"} <= set(data)
    assert data["atom"] == "H1"
    assert data["element"] == "H"
    # new key: raw value, full float precision
    assert repr(data["scaled_ppm"]) == repr(_RAW_SCALED)
    assert data["scaled_ppm"] != round(_RAW_SCALED, 4)
    # other numbers keep their pre-v2 round-4 serialization (unchanged keys)
    assert repr(data["exp_ppm"]) == repr(round(_EXP, 4))
    assert repr(data["calc_ppm"]) == repr(round(_CALC, 4))
    assert repr(data["residual"]) == repr(round(_RESIDUAL, 4))


# ---------------------------------------------------------------------------
# NmrReport.as_dict: schema v2 top level
# ---------------------------------------------------------------------------


def test_v2_top_level_schema_fields() -> None:
    data = _report().as_dict()
    assert data["schema_version"] == REPORT_SCHEMA_VERSION == 2
    # D-phase placeholder (todo 29 populates) — stays null until then
    assert data["protocol_id"] is None
    # every v1 top-level key is still present (additive schema only)
    assert {
        "summary",
        "candidates",
        "config",
        "error_model",
        "dp5_mode",
        "fchl_kernel",
        "note",
    } <= set(data)
    assert data["error_model"] == "goodman-legacy"
    assert data["dp5_mode"] == "fchl"
    assert data["fchl_kernel"] == "qml"


def test_config_block_is_full_effective_config() -> None:
    report = _report()
    data = report.as_dict()
    # schema v2: config IS the full T21 effective-config record
    assert data["config"] == report.config.to_dict()
    assert set(data["config"]) == {
        "nuclei",
        "nmr_method",
        "nmr_basis",
        "solvent",
        "solvent_model",
        "tms_shieldings",
        "tms_1h",
        "tms_13c",
        "boltzmann_temp",
        "energy_window_kcal",
        "max_conformers",
        "error_model",
        "conformer_preset",
        "strict_equivalence",
        "protocol_fingerprint",
    }
    # the old 3 keys remain readable for existing consumers
    for key in ("nmr_method", "nmr_basis", "solvent"):
        assert key in data["config"]
    assert data["config"]["nmr_method"] == "mPW1PW91"
    # protocol_fingerprint (D-phase, T29) stays null in the config record too
    assert data["config"]["protocol_fingerprint"] is None


def test_v1_reader_keys_preserved_per_candidate() -> None:
    """Frontend/API-consumed per-candidate keys survive schema v2 untouched."""
    data = _report().as_dict()
    cand = data["candidates"][0]
    assert {
        "index",
        "label",
        "dp4_probability",
        "dp5_probability",
        "dp5_diagnostic_score",
        "dp5_kernel",
        "evidence",
        "probability",
        "n_conformers",
        "regression",
        "assignment",
        "conformers",
    } <= set(cand)
    assert cand["assignment"][0].keys() >= {"atom", "element", "exp_ppm", "calc_ppm", "residual"}
    assert cand["conformers"] == [
        {"id": "conf_000", "boltzmann_weight": 0.5},
        {"id": "conf_001", "boltzmann_weight": 0.3},
        {"id": "conf_002", "boltzmann_weight": 0.14876},
    ]


# ---------------------------------------------------------------------------
# per-candidate analysis_status + coverage
# ---------------------------------------------------------------------------


def test_candidate_analysis_status_derivation() -> None:
    data = _report().as_dict()
    valid, excluded, legacy = data["candidates"]
    # from evidence.status when the gate ran
    assert valid["analysis_status"] == "valid"
    assert excluded["analysis_status"] == "invalid"
    # neither evidence nor probability block → unknown (None), never invented
    assert legacy["analysis_status"] is None


def test_candidate_analysis_status_falls_back_to_typed_dp4() -> None:
    """No evidence but a typed block → the typed DP4 status (brief contract)."""
    cr = CandidateResult(
        index=0,
        label="typed-only",
        probability=_probability(
            "evidence_insufficient", "unavailable", None, None, "not_evaluated", "not_evaluated"
        ),
    )
    assert cr.as_dict()["analysis_status"] == "evidence_insufficient"


def test_coverage_per_candidate() -> None:
    data = _report().as_dict()
    valid, excluded, legacy = data["candidates"]

    cov = valid["coverage"]
    assert cov["expected_signals"] == 4
    assert cov["matched_signals"] == 4
    assert cov["per_nucleus"] == {
        "13C": {"expected": 1, "matched": 1},
        "1H": {"expected": 3, "matched": 3},
    }
    # successful population = conformers present with complete shieldings
    assert cov["n_conformers_successful"] == 3
    assert cov["successful_population"] == pytest.approx(sum(_CONF_WEIGHTS))
    # the pre-GIAO selected set is NOT recorded — stays null, never invented
    assert cov["n_conformers_selected"] is None
    assert cov["selected_population"] is None

    cov = excluded["coverage"]
    assert cov["expected_signals"] == 4
    assert cov["matched_signals"] == 0
    assert cov["n_conformers_successful"] == 0
    assert cov["successful_population"] == 0.0  # measured zero, not a missing value
    assert cov["selected_population"] is None

    # no evidence at all → signal counts are null, NEVER coerced to 0
    cov = legacy["coverage"]
    assert cov["expected_signals"] is None
    assert cov["matched_signals"] is None
    assert cov["per_nucleus"] is None
    assert cov["successful_population"] == 0.0


def test_probability_state_block_preserved_with_nulls() -> None:
    data = json.loads(json.dumps(_report().as_dict()))
    excluded = data["candidates"][1]
    assert excluded["probability"]["dp4"]["probability"] is None
    assert excluded["probability"]["dp5"]["probability"] is None
    assert excluded["probability"]["dp5"]["status"] == "unavailable"
    assert excluded["dp4_probability"] is None
    assert excluded["dp5_probability"] is None


# ---------------------------------------------------------------------------
# calibration provenance
# ---------------------------------------------------------------------------


def test_provenance_block() -> None:
    data = _report().as_dict()
    prov = data["provenance"]
    assert prov["error_model"] == "goodman-legacy"
    assert prov["dp5_mode"] == "fchl"
    # candidates disagree → "mixed"; per-candidate values stay in probability blocks
    assert prov["calibration_status"] == {
        "dp4": "mixed",  # unvalidated_protocol vs not_evaluated
        "dp5": "mixed",  # goodman_kde vs not_evaluated
    }
    assert prov["tms_references"] == _report().config.tms_shieldings
    # default config values come from the Goodman TMSdata table
    assert prov["tms_source"] == "goodman_tmsdata"


def test_calibration_status_single_and_absent() -> None:
    one = NmrReport(
        candidates=[
            CandidateResult(
                index=0,
                label="only",
                probability=_probability(
                    "valid", "valid", 0.5, 0.4, "unvalidated_protocol", "goodman_kde"
                ),
            )
        ]
    )
    prov = one.as_dict()["provenance"]
    assert prov["calibration_status"] == {
        "dp4": "unvalidated_protocol",
        "dp5": "goodman_kde",
    }

    bare = NmrReport(candidates=[CandidateResult(index=0, label="bare")])
    prov = bare.as_dict()["provenance"]
    # no typed block anywhere → null, not "" and not 0
    assert prov["calibration_status"] == {"dp4": None, "dp5": None}


def test_tms_source_classifies_custom_and_unknown() -> None:
    custom = NmrReport(
        config=NmrConfig(tms_shieldings={"1H": 30.0, "13C": 180.0}),
    )
    assert custom.as_dict()["provenance"]["tms_source"] == "custom"

    unknown = NmrReport(
        config=NmrConfig(nmr_method="B3LYP", nmr_basis="def2-TZVP"),
    )
    assert unknown.as_dict()["provenance"]["tms_source"] == "unknown"


# ---------------------------------------------------------------------------
# v2 round-trip through the real writer
# ---------------------------------------------------------------------------


def test_write_json_report_v2_round_trip(tmp_path: Path) -> None:
    path = write_json_report(_report(), tmp_path / "nmr_report.json")
    text = path.read_text(encoding="utf-8")
    data = json.loads(text)

    assert data["schema_version"] == 2
    assert data["protocol_id"] is None
    assert data["config"] == _report().config.to_dict()
    assert "coverage" in data["candidates"][0]
    assert "analysis_status" in data["candidates"][0]
    assert "provenance" in data

    # nulls survive as JSON null — never coerced
    assert '"protocol_id": null' in text
    assert '"selected_population": null' in text
    assert '"n_conformers_selected": null' in text
    assert '"analysis_status": null' in text
    assert '"probability": null' in text
    assert data["candidates"][1]["probability"]["dp5"]["probability"] is None

    # raw scaled_ppm survives the file round-trip byte-exact
    assert repr(data["candidates"][0]["assignment"][0]["scaled_ppm"]) == repr(_RAW_SCALED)


# ---------------------------------------------------------------------------
# legacy (v1) compatibility: read path unchanged + unknown-note rendering
# ---------------------------------------------------------------------------


def test_legacy_v1_fixture_reads_and_renders_unknown_note() -> None:
    """A v1 payload (no schema_version) parses and renders the legacy note."""
    payload = json.loads(_LEGACY_V1_FIXTURE)  # reads without error
    assert "schema_version" not in payload
    # old fields stay readable exactly as before
    assert payload["config"]["nmr_method"] == "mPW1PW91"
    assert payload["candidates"][0]["assignment"][0]["exp_ppm"] == 4.0
    assert payload["summary"]["winner"]["dp4"] == 0.5
    # validation state is unknown for historical reports
    assert report_validation_note(payload) == LEGACY_REPORT_NOTE == "历史报告：验证状态未知"


def test_v2_payload_renders_no_unknown_note(tmp_path: Path) -> None:
    path = write_json_report(_report(), tmp_path / "nmr_report.json")
    data = json.loads(path.read_text(encoding="utf-8"))
    assert report_validation_note(data) is None


def test_report_validation_note_tolerates_foreign_payloads() -> None:
    # non-dict values and absent keys never raise — pure display classifier
    assert report_validation_note({}) == LEGACY_REPORT_NOTE
    assert report_validation_note({"schema_version": 1}) == LEGACY_REPORT_NOTE
    assert report_validation_note({"schema_version": 99}) == LEGACY_REPORT_NOTE
    assert report_validation_note({"schema_version": 2}) is None


# ---------------------------------------------------------------------------
# JSON ↔ XLSX numeric consistency (openpyxl read-back)
# ---------------------------------------------------------------------------


def test_json_and_xlsx_carry_same_numbers(tmp_path: Path) -> None:
    """Both artifacts serialize the SAME numbers from the same raw fields."""
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    report = _report()
    json_path = write_json_report(report, tmp_path / "nmr_report.json")
    xlsx_path = write_xlsx_report(report, tmp_path / "nmr_assignment.xlsx")
    assert xlsx_path is not None

    data = json.loads(json_path.read_text(encoding="utf-8"))
    wb = load_workbook(xlsx_path)

    for cand in data["candidates"]:
        ws = wb[f"cand_{cand['index']}"]
        values = [row for row in ws.iter_rows(values_only=True)]

        # assignment rows: same tuple on both sides, scaled_ppm raw in both
        for offset, a in enumerate(cand["assignment"], start=1):
            assert values[offset] == (
                a["atom"],
                a["element"],
                a["exp_ppm"],
                a["calc_ppm"],
                a["scaled_ppm"],
                a["residual"],
            )
        # raw precision visible: the XLSX cell is NOT the round-4 display value
        if cand["assignment"]:
            scaled_cell = values[1][4]
            assert repr(scaled_cell) == repr(_RAW_SCALED)

        # DP4 / DP5 rows match the JSON flat fields (round-6 serialization)
        dp4_row = next(r for r in values if r and r[0] == "DP4")
        dp5_row = next(r for r in values if r and r[0] == "DP5")
        assert dp4_row[1] == cand["dp4_probability"]
        assert dp5_row[1] == cand["dp5_probability"]

        # regression block matches the JSON regression record (round-6)
        for nucleus, reg in cand["regression"].items():
            start = next(i for i, r in enumerate(values) if r and r[0] == f"regression[{nucleus}]")
            assert values[start][2] == reg["slope"]
            assert values[start + 1][2] == reg["intercept"]
            assert values[start + 2][2] == reg["r_squared"]
            assert values[start + 3][2] == reg["mae"]
